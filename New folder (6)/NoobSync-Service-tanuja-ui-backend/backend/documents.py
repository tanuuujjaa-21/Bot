"""
documents.py
Turns an uploaded file (PDF, Word, text, Markdown, CSV) into plain text for
the RAG layer, and keeps filenames safe.
"""

import csv
import io
import os
import re
import unicodedata

from backend.config import ALLOWED_EXTENSIONS


class DocumentError(Exception):
    """Message is safe to show to the user."""


def extension_of(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def display_name(raw: str) -> str:
    """A safe, readable filename for the UI. The original is never used as a
    path on disk (files are stored under generated ids)."""
    name = os.path.basename((raw or "").replace("\\", "/"))
    name = "".join(ch for ch in unicodedata.normalize("NFKC", name)
                   if unicodedata.category(ch)[0] != "C")
    name = re.sub(r"\s+", " ", name).strip()
    return (name or "document")[:120]


def check_extension(filename: str) -> str:
    ext = extension_of(filename)
    if ext not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(e.lstrip(".").upper() for e in ALLOWED_EXTENSIONS)
        raise DocumentError(f"Unsupported file type. Upload one of: {allowed}.")
    return ext


def _decode(data: bytes) -> str:
    # UTF-16 only when a byte-order mark says so: without one, almost any
    # even-length byte string "decodes" as UTF-16 into garbage.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeError:
            pass
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise DocumentError("This PDF is password-protected. Remove the password and try again.")
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("This PDF could not be read. It may be damaged.") from exc


def _from_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("This Word file could not be read. It may be damaged.") from exc

    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def _from_csv(data: bytes) -> str:
    """Each row becomes 'Header: value; Header: value' so a row stays
    meaningful on its own when it lands in a retrieved chunk."""
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return ""
    header, body = rows[0], rows[1:]
    if not body:
        return ", ".join(header)
    lines = []
    for r in body:
        lines.append("; ".join(
            f"{(header[i] if i < len(header) else f'Column {i + 1}').strip()}: {v.strip()}"
            for i, v in enumerate(r) if v.strip()
        ))
    return "\n".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    ext = check_extension(filename)
    if ext == ".pdf":
        text = _from_pdf(data)
    elif ext == ".docx":
        text = _from_docx(data)
    elif ext == ".csv":
        text = _from_csv(data)
    else:  # .txt, .md
        text = _decode(data)

    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentError(
            "No readable text was found in this file. "
            "Scanned or image-only documents are not supported."
        )
    return text
