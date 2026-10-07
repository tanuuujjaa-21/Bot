"""Backend checks that need no model, network or API key.
Run from the project root:  python tests/test_backend.py   (or: pytest tests)
"""
import os, sys, tempfile, io
from pathlib import Path

os.environ["DATA_DIR"] = tempfile.mkdtemp()  # keep tests away from real data/
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


def test_backend():
    _run()


def _run():
    from backend import db, auth, documents
    from backend.documents import DocumentError
    db.init_db()

    # --- validation
    assert auth.clean_name("  Tanuja   Gaddam ") == "Tanuja Gaddam"
    assert auth.clean_name("Anne-Marie O'Neil") == "Anne-Marie O'Neil"
    for bad in ["", "A", "1234", "Bob<script>", "x"*61, "!!"]:
        try: auth.clean_name(bad); raise AssertionError(f"name accepted: {bad!r}")
        except auth.RegistrationError: pass
    assert auth.clean_phone("+91 98765-43210") == "+919876543210"
    assert auth.clean_phone("(987) 654 3210") == "9876543210"
    for bad in ["", "12345", "abcdefghij", "+1234567890123456", "98765 4321a"]:
        try: auth.clean_phone(bad); raise AssertionError(f"phone accepted: {bad!r}")
        except auth.RegistrationError: pass

    # --- users, tokens, isolation
    u1 = db.create_user("Asha", "9876543210"); u2 = db.create_user("Ravi", "9876543210")
    assert db.phone_seen_before("9876543210") and not db.phone_seen_before("9000000000")
    t1, t2 = auth.issue_token(u1["id"]), auth.issue_token(u2["id"])
    assert auth.current_user(f"Bearer {t1}")["id"] == u1["id"]
    try: auth.current_user("Bearer nope"); raise AssertionError("bad token accepted")
    except Exception as e: assert e.status_code == 401
    try: auth.current_user(None); raise AssertionError("no header accepted")
    except Exception as e: assert e.status_code == 401
    auth.revoke_token(t1)
    try: auth.current_user(f"Bearer {t1}"); raise AssertionError("revoked token still valid")
    except Exception as e: assert e.status_code == 401
    assert db.get_user(u1["id"])["name"] == "Asha"

    # --- history + ownership
    c = db.create_conversation(u2["id"])
    assert db.list_conversations(u2["id"]) == []                       # empty chats are hidden
    assert db.get_conversation(u1["id"], c["id"]) is None            # other user can't see it
    assert db.get_conversation(u2["id"], c["id"])["title"] == "New chat"
    q, a = db.add_exchange(c["id"], "What is ConnectOps?", "Zoho & WhatsApp automation.", '{"type":"rag_pipeline","files":[]}', 120.5, "What is ConnectOps?")
    db.add_exchange(c["id"], "second", "answer2", None, 1.0, "should not replace title")
    assert [x["id"] for x in db.list_conversations(u2["id"])] == [c["id"]]
    assert db.list_conversations(u1["id"]) == []
    msgs = db.list_messages(c["id"])
    assert [m["role"] for m in msgs] == ["user","assistant","user","assistant"]
    assert db.get_conversation(u2["id"], c["id"])["title"] == "What is ConnectOps?"
    assert not db.delete_conversation(u1["id"], c["id"])              # not owner
    assert db.delete_conversation(u2["id"], c["id"])
    assert db.list_messages(c["id"]) == []                            # cascade

    # --- documents: ownership + restart behaviour
    d = db.add_document("d1", u2["id"], "a.pdf", ".pdf", 10, "/x/d1.txt")
    assert d["status"] == "processing" and db.get_document(u1["id"], "d1") is None
    db.init_db()                                                      # simulated restart
    assert db.get_document(u2["id"], "d1")["status"] == "error"
    assert db.delete_document(u1["id"], "d1") is None and db.delete_document(u2["id"], "d1")

    # --- rate limit
    import fastapi
    for i in range(10): auth.check_register_rate("1.2.3.4")
    try: auth.check_register_rate("1.2.3.4"); raise AssertionError("not limited")
    except Exception as e: assert e.status_code == 429
    auth.check_register_rate("5.6.7.8")

    # --- extraction for each supported type
    assert "hello" in documents.extract_text("n.txt", b"hello world")
    assert "# Title" in documents.extract_text("n.md", "# Title\nbody".encode())
    assert "Name: Asha; City: Pune" in documents.extract_text("n.csv", b"Name,City\nAsha,Pune\nRavi,Delhi\n")
    assert "caf\u00e9" in documents.extract_text("n.txt", "caf\u00e9".encode("cp1252"))
    import docx
    dd = docx.Document(); dd.add_paragraph("ConnectOps handles Zoho."); t = dd.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text="Plan"; t.rows[0].cells[1].text="SME"
    b = io.BytesIO(); dd.save(b)
    out = documents.extract_text("n.docx", b.getvalue()); assert "ConnectOps handles Zoho." in out and "Plan | SME" in out
    from pypdf import PdfWriter
    w = PdfWriter(); w.add_blank_page(72,72); pb = io.BytesIO(); w.write(pb)
    for fn, data, why in [("blank.pdf", pb.getvalue(), "No readable text"), ("junk.pdf", b"not a pdf", "could not be read"),
                          ("e.txt", b"   \n ", "No readable text"), ("bad.docx", b"zzz", "could not be read"),
                          ("x.exe", b"MZ", "Unsupported"), ("noext", b"hi", "Unsupported")]:
        try: documents.extract_text(fn, data); raise AssertionError(f"{fn} accepted")
        except DocumentError as e: assert why in str(e), (fn, str(e))
    assert documents.display_name("../../etc/pass\x00wd.txt") == "passwd.txt"
    assert documents.display_name("C:\\Users\\a\\report.pdf") == "report.pdf"
    print("ALL BACKEND CHECKS PASSED")


if __name__ == "__main__":
    _run()
